__import__('pysqlite3')
import sys
sys.modules['sqlite3'] = sys.modules.pop('pysqlite3')
import streamlit as st
from langchain_core.messages import AIMessage, HumanMessage
from langchain_community.document_loaders import WebBaseLoader, PyMuPDFLoader, Docx2txtLoader, UnstructuredExcelLoader, UnstructuredFileLoader, UnstructuredPDFLoader
from langchain_groq import ChatGroq
from langchain.schema.document import Document
from langchain_google_genai import GoogleGenerativeAIEmbeddings
from langchain.text_splitter import RecursiveCharacterTextSplitter
# from langchain_community.vectorstores import Chroma
from langchain_chroma import Chroma
from dotenv import load_dotenv
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain.chains import create_history_aware_retriever, create_retrieval_chain
from langchain.chains.combine_documents import create_stuff_documents_chain
import time
import pandas as pd
import pytesseract
import openpyxl
import tempfile
import os
import nltk

#Added the required dwonloads from nltk to run the unstructured loader
try:
    nltk.download('punkt_tab')
    nltk.download('averaged_perceptron_tagger_eng')
except:
    pass

load_dotenv()

# pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

def authenticate(username: str, password: str):
    
    if username == "HR" and password == "1234":
        return ["HR", "Legal", "Commercial", "Domain 5", "Domain 6", "Domain 7"]
    
    elif username == "IT" and password == "1234":
        return ["IT", "Domain 8", "Domain 9", "Domain 10"]
    
    elif username == "admin" and password == "1234":
        return ["HR", "IT", "Legal", "Commercial", "Domain 5", "Domain 6", "Domain 7", "Domain 8", "Domain 9", "Domain 10"]
    
    else:
        return None
    
    
def create_vectorstore_from_pdf(docs, domain):
    # get the text in document form
    documents = []
    for doc in docs:
        if os.path.splitext(doc)[1] == ".pdf":
            loader = PyMuPDFLoader(doc)
            document = loader.load()
            # if document[0].page_content != "" or document[0].page_content != " " or document[0].page_content != None:
            # st.write(f"PymuPDF: {document}")
            # for d in document:
            #     documents.append(d)   
                    
            if document[0].page_content == "" or document[0].page_content == " " or document[0].page_content == None:
                start_time = time.time()
                st.write("Taking UnStructured Route")
                loader = UnstructuredFileLoader(doc)
                print(loader)
                st.write("Started Extracting")
                document = loader.load()
                print(document)
                st.write("extraction done.")
                end_time = time.time()
                calculated_time = end_time - start_time
                print("Time to process the Unstructured PDF", calculated_time)
                
                # st.write(document)
            for d in document:
                # st.write(d)
                documents.append(d)
                

        elif os.path.splitext(doc)[1] == ".docx":
            loader = Docx2txtLoader(doc)
            document = loader.load()
            for d in document:
                documents.append(d)
                
        
        elif os.path.splitext(doc)[1] == ".xlsx":
            # loader = UnstructuredExcelLoader(doc)
            # document = loader.load()
            # for d in document:
            #     documents.append(d)
            workbook = openpyxl.load_workbook(doc)

            sheet = workbook.active  # You can specify a sheet name if needed: workbook['SheetName']

            # Extract data and convert it to text
            extracted_text = ""

            # Loop through rows and columns
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                extracted_text += row_text + "\n"
            
            document = Document(page_content=extracted_text, metadata = {"source": doc})
            documents.append(document)
        
    # st.write(documents)
    # split the document into chunks
    text_splitter = RecursiveCharacterTextSplitter()
    document_chunks = text_splitter.split_documents(documents)
    # st.write(document_chunks)
    # create a vectorstore from the chunks
    vector_store = Chroma(persist_directory=f"{domain}/", embedding_function=GoogleGenerativeAIEmbeddings(model="models/embedding-001"))
    # vector_store = Chroma.from_documents(document_chunks, GoogleGenerativeAIEmbeddings(model="models/embedding-001"))
    vector_store.add_documents(document_chunks)
    

    
st.set_page_config(page_title='chat with your Documents', page_icon=':books:')
st.image("LOGO.png", width=455)
st.header("Upload the Documents below")


    
# try:
if st.session_state.domains:
    
    domain = st.sidebar.selectbox("Select a Domain", options=st.session_state.domains)
    
    uploaded_files = st.file_uploader("uplod your pdf here", type=["pdf", "docx", "xlsx"], accept_multiple_files=True)

    if uploaded_files:
        uploaded_file_names = [uploaded_file.name for uploaded_file in uploaded_files]
        if st.button("Process"):
            temp_file_paths = []
            for uploaded_file in uploaded_files:
                print(uploaded_file)
                temp_dir = tempfile.gettempdir()  ##Get the system's temporary directory
                
                temp_file_path = os.path.join(temp_dir, uploaded_file.name)
                with open(temp_file_path, "wb") as temp_file:
                    temp_file.write(uploaded_file.getbuffer())
                temp_file_paths.append(temp_file_path)
                
            create_vectorstore_from_pdf(temp_file_paths, domain)
            documents_dataframe = pd.DataFrame(columns=["Domain", "Document"])
            rows = []
            for uploaded_file_name in uploaded_file_names:
                rows.append({"Domain": domain, "Document": uploaded_file_name})
            documents_dataframe = pd.concat([documents_dataframe, pd.DataFrame(rows)], ignore_index=True)
            print(documents_dataframe)
            csv_file_path = "Document_Data.csv"
            if not os.path.isfile(csv_file_path):
                documents_dataframe.to_csv(csv_file_path, index=False)
            else:
                documents_dataframe.to_csv(csv_file_path, mode='a', header=False, index=False)
                
            ##Showing the updated documents
            st.subheader("Uploaded Documents")
            documents = pd.read_csv(csv_file_path)
            st.dataframe(documents)
                
            st.success("Documents Successfully Store. You can now chat with the documents.")
# except:
#     st.warning("Please login first.")
